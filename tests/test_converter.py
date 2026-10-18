"""пять примеров: входной файл Markdown и бинарный Docx в stdout."""

from io import BytesIO
from pathlib import Path
import subprocess
import sys

import pytest
from docx import Document


@pytest.mark.parametrize('source', sorted(Path('examples').glob('*.md')))
def test_examples(source, tmp_path):
    """проверяет читаемость stdout как docx, без служебного текста."""
    result = subprocess.run(
        [sys.executable, '-m', 'src.converter', str(source)],
        capture_output=True, check=True,
    )
    document = Document(BytesIO(result.stdout))
    output = tmp_path / f'{source.stem}.docx'
    output.write_bytes(result.stdout)
    assert any(paragraph.text for paragraph in document.paragraphs)
    assert Document(output).paragraphs[0].text == document.paragraphs[0].text
    assert result.stderr == b''

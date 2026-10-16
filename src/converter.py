"""чтение Markdown из stdin и сохранение документа на диск."""

import argparse
import sys
from pathlib import Path

from docx import Document
from markdown_it import MarkdownIt


HEADING_OFFSET = 1


def add_inline(paragraph, tokens):
    """переносит текст и выделение из токенов Markdown в абзац.

    стек хранит открытые выделения, поэтому жирный и курсив можно
    вкладывать друг в друга. ссылки сохраняются вместе с адресом.
    """
    styles = []
    for token in tokens:
        if token.type.endswith('_open'):
            styles.append(token.type)
        elif token.type.endswith('_close'):
            if styles:
                styles.pop()
        elif token.type in ('softbreak', 'hardbreak'):
            paragraph.add_run().add_break()
        else:
            run = paragraph.add_run(token.content)
            run.bold = 'strong_open' in styles
            run.italic = 'em_open' in styles
            if token.type == 'code_inline':
                run.font.name = 'Courier New'
        if token.type == 'link_open':
            paragraph.add_run(f" ({token.attrGet('href')}) ")


def add_block(document, token, lists):
    """создаёт абзац нужного вида для очередного блока Markdown."""
    if token.type == 'heading_open':
        level = int(token.tag[HEADING_OFFSET:])
        return document.add_heading(level=level)
    if token.type == 'paragraph_open':
        style = None
        if lists:
            style = ('List Number' if lists[-1] == 'ordered_list_open'
                     else 'List Bullet')
        return document.add_paragraph(style=style)
    if token.type in ('fence', 'code_block'):
        run = document.add_paragraph().add_run(token.content)
        run.font.name = 'Courier New'
    if token.type == 'hr':
        document.add_paragraph('────────────────')
    return None


def convert(text, output):
    """разбирает Markdown и сохраняет заголовки, списки, текст и код.

    markdown-it отвечает за синтаксический разбор, а python-docx
    создаёт настоящий документ Office Open XML с расширением .docx.
    """
    document = Document()
    tokens = MarkdownIt('commonmark').parse(text)
    paragraph = None
    lists = []
    for token in tokens:
        if token.type in ('bullet_list_open', 'ordered_list_open'):
            lists.append(token.type)
        elif token.type in ('bullet_list_close', 'ordered_list_close'):
            lists.pop()
        elif token.type == 'inline':
            add_inline(paragraph, token.children)
        else:
            block = add_block(document, token, lists)
            if block is not None:
                paragraph = block
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main():
    """получает путь результата из аргумента, а Markdown — из stdin."""
    parser = argparse.ArgumentParser(description='Markdown -> Docx')
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    convert(sys.stdin.read(), args.output)
    print(args.output)


if __name__ == '__main__':
    main()
